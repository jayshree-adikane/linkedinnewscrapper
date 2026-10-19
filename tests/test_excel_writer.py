from datetime import datetime

from openpyxl import load_workbook

from excel.excel_writer import errors_workbook_bytes, workbook_bytes, write_workbook
from scraper.models import (
    STATUS_FAILED,
    STATUS_INCOMPLETE,
    STATUS_SKIPPED,
    STATUS_SUCCESS,
    ProspectRecord,
    RunInfo,
)

URL = "https://www.linkedin.com/in/a/"


def sample():
    return [
        ProspectRecord(2, URL, full_name="=1+1", current_title="Engineer", status=STATUS_SUCCESS),
        ProspectRecord(3, "https://www.linkedin.com/in/b/", status=STATUS_INCOMPLETE, error_category="incomplete_data", error_message="Not available on page: location"),
        ProspectRecord(4, "https://www.linkedin.com/in/c/", status=STATUS_FAILED, error_category="captcha", error_message="CAPTCHA"),
        ProspectRecord(5, "", status=STATUS_SKIPPED, error_category="invalid_url", error_message="bad"),
    ]


def info():
    return RunInfo(datetime(2026, 1, 1, 9, 0), datetime(2026, 1, 1, 9, 5), 4, 3)


def test_workbook_sheets_and_content(tmp_path):
    path = tmp_path / "out.xlsx"
    write_workbook(path, sample(), info())
    wb = load_workbook(path)
    assert wb.sheetnames == ["Prospects", "Errors", "Summary"]

    prospects = wb["Prospects"]
    assert prospects.max_row == 4  # header + success + incomplete + failed (skipped excluded)
    headers = [c.value for c in prospects[1]]
    assert "Current Job Start Month" in headers and "Previous Job End (Month Year)" in headers
    name_cell = prospects.cell(row=2, column=headers.index("Full Name") + 1)
    assert name_cell.value == "=1+1" and name_cell.data_type == "s"  # not a formula
    url_cell = prospects.cell(row=2, column=headers.index("LinkedIn Profile URL") + 1)
    assert url_cell.hyperlink is not None

    errors = wb["Errors"]
    assert [c.value for c in errors[1]] == ["Profile URL", "Original Row", "Processing Status", "Error Category", "Diagnostic Message"]
    assert errors.max_row == 3  # failed + skipped

    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Total input records"] == 4
    assert summary["Unique profile URLs"] == 3
    assert summary["Successfully processed records"] == 1
    assert summary["Records with incomplete data"] == 1
    assert summary["Failed records"] == 1
    assert summary["Skipped records"] == 1
    assert summary["Run start time"] == "2026-01-01 09:00:00"


def test_bytes_helpers_roundtrip(tmp_path):
    import io

    assert load_workbook(io.BytesIO(workbook_bytes(sample(), info()))).sheetnames[0] == "Prospects"
    assert load_workbook(io.BytesIO(errors_workbook_bytes(sample()))).sheetnames == ["Errors"]
