"""Build the output workbook (Prospects / Errors / Summary)."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional, Union

import pandas as pd
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from scraper.models import (
    PROSPECT_COLUMNS,
    STATUS_FAILED,
    STATUS_INCOMPLETE,
    STATUS_SKIPPED,
    STATUS_SUCCESS,
    ProspectRecord,
    RunInfo,
)

ERROR_COLUMNS = [
    "Profile URL",
    "Original Row",
    "Processing Status",
    "Error Category",
    "Diagnostic Message",
]
_URL_HEADERS = {"LinkedIn Profile URL", "Profile URL"}
Target = Union[str, Path, io.BytesIO]


def _fmt(value: Optional[datetime]) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value else ""


def prospects_frame(results: Iterable[ProspectRecord]) -> pd.DataFrame:
    """Processed profiles (skipped rows are reported on the Errors sheet only)."""
    rows = [r.to_row() for r in results if r.status != STATUS_SKIPPED]
    return pd.DataFrame(rows, columns=[h for h, _ in PROSPECT_COLUMNS])


def errors_frame(results: Iterable[ProspectRecord]) -> pd.DataFrame:
    """Failed and skipped rows with a short diagnostic."""
    rows = [
        {
            "Profile URL": r.url,
            "Original Row": r.row_number,
            "Processing Status": r.status,
            "Error Category": r.error_category,
            "Diagnostic Message": r.error_message,
        }
        for r in results
        if r.status in (STATUS_FAILED, STATUS_SKIPPED)
    ]
    return pd.DataFrame(rows, columns=ERROR_COLUMNS)


def summary_frame(results: list[ProspectRecord], info: RunInfo) -> pd.DataFrame:
    """Key/value summary of the run."""

    def count(status: str) -> int:
        return sum(r.status == status for r in results)

    rows = [
        ("Total input records", info.total_input),
        ("Unique profile URLs", info.unique_urls),
        ("Successfully processed records", count(STATUS_SUCCESS)),
        ("Records with incomplete data", count(STATUS_INCOMPLETE)),
        ("Failed records", count(STATUS_FAILED)),
        ("Skipped records", count(STATUS_SKIPPED)),
        ("Run start time", _fmt(info.started_at)),
        ("Run end time", _fmt(info.finished_at)),
    ]
    if info.stop_reason:
        rows.append(("Run stopped early because", info.stop_reason))
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def _format_sheet(ws: Worksheet) -> None:
    """Text formatting, hyperlinks, header style, widths, formula-injection guard."""
    headers = [c.value for c in ws[1]]
    url_cols = {i for i, h in enumerate(headers) if h in _URL_HEADERS}
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for idx, cell in enumerate(row):
            if isinstance(cell.value, str):
                cell.data_type = "s"  # never let "=..." become a formula
                cell.number_format = "@"
                if idx in url_cols and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    for column in ws.columns:
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(longest + 2, 12), 60)


def _write(target: Target, frames: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            _format_sheet(writer.sheets[name])


def write_workbook(target: Target, results: list[ProspectRecord], info: RunInfo) -> None:
    """Write the full 3-sheet workbook to a path or file-like object."""
    _write(
        target,
        {
            "Prospects": prospects_frame(results),
            "Errors": errors_frame(results),
            "Summary": summary_frame(results, info),
        },
    )


def workbook_bytes(results: list[ProspectRecord], info: RunInfo) -> bytes:
    """The full workbook as bytes (for Streamlit downloads)."""
    buffer = io.BytesIO()
    write_workbook(buffer, results, info)
    return buffer.getvalue()


def errors_workbook_bytes(results: list[ProspectRecord]) -> bytes:
    """Stand-alone error report."""
    buffer = io.BytesIO()
    _write(buffer, {"Errors": errors_frame(results)})
    return buffer.getvalue()


def write_errors_workbook(path: Path, results: list[ProspectRecord]) -> None:
    path.write_bytes(errors_workbook_bytes(results))
